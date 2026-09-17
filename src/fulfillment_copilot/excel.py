"""Excel 交付物：把分析结果落成业务同学可直接使用的多 Sheet 工作簿。

对应 JD「熟练使用 Excel / 具备数据分析与报告撰写能力」——但更进一步：
不是人肉整理 Excel，而是**让分析结果自动生成一份带格式、带条件格式、可直接下发的 Excel**，
把运营同学的时间从"整理表格"转移到"解决问题"。

依赖 ``openpyxl``（可选）：未安装时 ``write_workbook`` 会抛出 ``ExcelUnavailable``，
调用方（CLI）据此提示 ``pip install openpyxl`` 并继续产出 Markdown/HTML 报告，
不会让整条流水线失败。
"""

from __future__ import annotations

from datetime import datetime
from typing import Any, Optional, Sequence

from .metrics import CATEGORY_LABELS, METRICS, format_metric

try:  # pragma: no cover - 取决于运行环境
    from openpyxl import Workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    OPENPYXL_AVAILABLE = True
except ImportError:  # pragma: no cover
    OPENPYXL_AVAILABLE = False


class ExcelUnavailable(RuntimeError):
    """未安装 openpyxl。"""


#: 表头样式
_HEADER_FILL = "1F4E79"
_HEADER_FONT = "FFFFFF"
_ACCENT_FILL = "DDEBF7"
_WARN_FILL = "FCE4D6"

#: 百分比列（用于条件格式与数字格式）
_RATE_KEYS = {
    "dispatch_ontime_rate",
    "arrive_ontime_rate",
    "finish_ontime_rate",
    "sla_achieve_rate",
    "first_fix_rate",
    "rework_rate",
    "reschedule_rate",
    "cancel_rate",
    "complaint_rate",
    "good_review_rate",
    "bad_review_rate",
    "gross_margin",
}

#: 越低越好的百分比列（条件格式方向相反）
_LOWER_BETTER = {"rework_rate", "reschedule_rate", "cancel_rate", "complaint_rate", "bad_review_rate"}


def excel_available() -> bool:
    return OPENPYXL_AVAILABLE


def _style_header(worksheet, columns: int, row: int = 1) -> None:
    for index in range(1, columns + 1):
        cell = worksheet.cell(row=row, column=index)
        cell.fill = PatternFill("solid", fgColor=_HEADER_FILL)
        cell.font = Font(color=_HEADER_FONT, bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.freeze_panes = worksheet.cell(row=row + 1, column=1)


def _autosize(worksheet, min_width: int = 10, max_width: int = 42) -> None:
    for column_cells in worksheet.columns:
        length = max((len(str(cell.value)) if cell.value is not None else 0) for cell in column_cells)
        width = min(max(length + 4, min_width), max_width)
        worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def _add_table(worksheet, headers: Sequence[str], rows: Sequence[Sequence[Any]], number_formats: Optional[dict[int, str]] = None) -> None:
    worksheet.append(list(headers))
    for row in rows:
        worksheet.append(list(row))
    _style_header(worksheet, len(headers))
    if number_formats:
        for column_index, fmt in number_formats.items():
            for row_index in range(2, worksheet.max_row + 1):
                worksheet.cell(row=row_index, column=column_index).number_format = fmt
    _autosize(worksheet)


def _percent_columns(keys: Sequence[str], offset: int) -> dict[int, str]:
    """把指标列映射为 Excel 列号 → 百分比/数值格式。"""
    formats: dict[int, str] = {}
    for index, key in enumerate(keys):
        if key in _RATE_KEYS:
            formats[index + offset] = "0.0%"
        elif METRICS.get(key) and METRICS[key].unit == "小时":
            formats[index + offset] = "0.0"
        elif METRICS.get(key) and METRICS[key].unit == "元":
            formats[index + offset] = "#,##0.0"
    return formats


def _dimension_sheet(workbook, title: str, dimension_label: str, rows: Sequence[dict], keys: Sequence[str]) -> None:
    worksheet = workbook.create_sheet(title)
    headers = [dimension_label, "工单量"] + [METRICS[key].name for key in keys]
    body = []
    for row in rows:
        body.append([row.get("label") or row.get("value"), row.get("orders")] + [row.get(key) for key in keys])
    formats = _percent_columns(keys, offset=3)
    formats[2] = "#,##0"
    _add_table(worksheet, headers, body, formats)
    _conditional_formats(worksheet, keys, first_data_row=2, first_metric_column=3, last_row=worksheet.max_row)


def _conditional_formats(worksheet, keys: Sequence[str], first_data_row: int, first_metric_column: int, last_row: int) -> None:
    """给比率列加三色阶：绿色好、红色差，一眼看出问题站点/品类。"""
    if last_row < first_data_row:
        return
    for index, key in enumerate(keys):
        if key not in _RATE_KEYS:
            continue
        column = get_column_letter(first_metric_column + index)
        cell_range = f"{column}{first_data_row}:{column}{last_row}"
        if key in _LOWER_BETTER:
            rule = ColorScaleRule(start_type="num", start_value=0, start_color="63BE7B",
                                  mid_type="percentile", mid_value=50, mid_color="FFEB84",
                                  end_type="max", end_color="F8696B")
        else:
            rule = ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                                  mid_type="percentile", mid_value=50, mid_color="FFEB84",
                                  end_type="max", end_color="63BE7B")
        worksheet.conditional_formatting.add(cell_range, rule)


def _overview_sheet(workbook, evidence: dict[str, Any]) -> None:
    worksheet = workbook.active
    worksheet.title = "总览"
    meta = evidence["meta"]
    overview = evidence["overview"]
    scorecard = evidence["scorecard"]

    worksheet.append([f"服务履约质量分析 · 数据截至 {meta['anchor'][:10]}"])
    worksheet["A1"].font = Font(bold=True, size=14, color=_HEADER_FILL)
    worksheet.append([f"覆盖最近 {meta['window_days']} 天 · {meta['orders']} 单 · {meta['sites']} 个站点 · 生成时间 {meta['generated_at']}"])
    worksheet.append([])
    worksheet.append(["履约质量综合得分", scorecard.get("total"), f"{scorecard.get('grade')} 级"])
    worksheet.append([])

    headers = ["指标", "类别", "当前值", "单位", "优化方向"]
    rows = []
    for key, item in overview.items():
        spec = METRICS[key]
        rows.append([spec.name, CATEGORY_LABELS.get(spec.category, spec.category), item.get("value"), spec.unit,
                     "越大越好" if spec.higher_is_better else "越小越好"])
    start_row = worksheet.max_row + 1
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    _style_header(worksheet, len(headers), row=start_row)
    formats = {}
    for index, key in enumerate(overview.keys()):
        if key in _RATE_KEYS:
            formats[index + 3] = "0.0%"
    for column_index, fmt in formats.items():
        for row_index in range(start_row + 1, worksheet.max_row + 1):
            worksheet.cell(row=row_index, column=column_index).number_format = fmt
    _autosize(worksheet)


def _trend_sheet(workbook, evidence: dict[str, Any]) -> None:
    worksheet = workbook.create_sheet("周趋势")
    keys = ["orders", "arrive_ontime_rate", "sla_achieve_rate", "dispatch_p90", "first_fix_rate", "complaint_rate"]
    headers = ["周期", "开始", "结束"] + ["工单量" if key == "orders" else METRICS[key].name for key in keys]
    body = []
    for row in evidence["trend"]:
        values = [row.get("orders") if key == "orders" else row.get(key) for key in keys]
        body.append([row.get("period"), row.get("start"), row.get("end")] + values)
    formats = {}
    for index, key in enumerate(keys):
        if key in _RATE_KEYS:
            formats[index + 4] = "0.0%"
        if key == "orders":
            formats[index + 4] = "#,##0"
    _add_table(worksheet, headers, body, formats)


def _anomaly_sheet(workbook, evidence: dict[str, Any]) -> None:
    worksheet = workbook.create_sheet("指标异常")
    headers = ["严重度", "维度", "对象", "指标", "基线", "当前值", "偏离度", "稳健z", "近期工单量", "建议动作"]
    body = []
    for item in evidence["anomalies"]["dimension"]:
        body.append(
            [
                item.get("severity"),
                item.get("dimension"),
                item.get("label") or item.get("value"),
                item.get("metric_name"),
                item.get("baseline_value"),
                item.get("current_value"),
                item.get("deviation"),
                item.get("robust_z"),
                item.get("recent_orders"),
                item.get("suggestion"),
            ]
        )
    _add_table(worksheet, headers, body, {7: "0.0%", 9: "#,##0"})
    if worksheet.max_row > 1:
        worksheet.conditional_formatting.add(
            f"A2:A{worksheet.max_row}",
            ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                           mid_type="num", mid_value=1, mid_color="FFEB84",
                           end_type="num", end_value=2, end_color="63BE7B"),
        )
        for row_index in range(2, worksheet.max_row + 1):
            if worksheet.cell(row=row_index, column=1).value == "高":
                for column_index in range(1, len(headers) + 1):
                    worksheet.cell(row=row_index, column=column_index).fill = PatternFill("solid", fgColor=_WARN_FILL)


def _rule_detail_sheet(workbook, evidence: dict[str, Any]) -> None:
    worksheet = workbook.create_sheet("需跟进工单")
    detail = evidence["anomalies"]["rule"].get("detail") or []
    headers = ["工单号", "异常类型", "严重度", "品类", "站点", "城市", "工程师", "下单时间", "具体问题"]
    body = [
        [item.get("order_id"), item.get("anomaly_type"), item.get("severity"), item.get("category"),
         item.get("site_name"), item.get("city"), item.get("engineer_id"), item.get("order_time"), item.get("detail")]
        for item in detail
    ]
    _add_table(worksheet, headers, body)
    worksheet.append([])
    worksheet.append(["说明：以上为规则引擎筛选出的候选工单明细（按严重度排序，最多 60 条）。"])
    worksheet.cell(row=worksheet.max_row, column=1).font = Font(italic=True, size=9, color="808080")


def _feedback_sheet(workbook, evidence: dict[str, Any]) -> None:
    worksheet = workbook.create_sheet("用户声音")
    feedback = evidence["feedback"]
    worksheet.append(["主题", "条数", "占负面比", "典型原声"])
    for theme in feedback.get("themes") or []:
        examples = (feedback.get("examples") or {}).get(theme["theme"]) or []
        worksheet.append([theme["theme"], theme["count"], theme["share"], examples[0] if examples else ""])
    _style_header(worksheet, 4)
    for row_index in range(2, worksheet.max_row + 1):
        worksheet.cell(row=row_index, column=3).number_format = "0.0%"
    worksheet.append([])
    worksheet.append(["高频词", "、".join(f"{word}({count})" for word, count in (feedback.get("keywords") or [])[:12])])
    _autosize(worksheet, max_width=60)


def _dispatch_sheet(workbook, evidence: dict[str, Any]) -> None:
    view = evidence.get("dispatch") or {}
    worksheet = workbook.create_sheet("派单评分卡")
    headers = ["站点", "类型", "城市", "工程师", "工单量", "单/人/天", "综合得分", "上门准时率", "一次完工率", "客诉率", "派单建议"]
    body = []
    for row in view.get("sites") or []:
        body.append([
            row.get("site_name"), row.get("site_type"), row.get("city"), row.get("engineers"), row.get("orders"),
            row.get("orders_per_engineer_day"), row.get("score"),
            row.get("arrive_ontime_rate"), row.get("first_fix_rate"), row.get("complaint_rate"), row.get("recommendation"),
        ])
    _add_table(worksheet, headers, body, {8: "0.0%", 9: "0.0%", 10: "0.0%"})
    worksheet.append([])
    worksheet.append(["建议规则"])
    for rule in view.get("rules") or []:
        worksheet.append([f"P{rule.get('priority')}", rule.get("name"), rule.get("trigger"), rule.get("action"), rule.get("evidence")])
    _autosize(worksheet, max_width=60)


def _notes_sheet(workbook, evidence: dict[str, Any]) -> None:
    worksheet = workbook.create_sheet("口径说明")
    notes = [
        ["项目", "说明"],
        ["数据来源", "合成数据（固定随机种子，可完全复现）；真实业务数据涉及用户隐私与商业机密，不进入公开仓库"],
        ["覆盖链路", "下单 → 预约受理 → 约定上门时段 → 派单 → 上门 → 完工 → 评价"],
        ["上门准时率", "实际上门时间 ≤ 约定上门时段 + 品类到达窗口"],
        ["一次完工率", "首次上门即解决、无返工的完工工单占比"],
        ["全链路SLA达成率", "下单到完工时长 ≤ 品类时效承诺的工单占比"],
        ["客诉率", "产生客诉的工单占比（客诉主题由客服系统归类）"],
        ["差评率", "评分 ≤ 2 的工单占已评价工单的比例"],
        ["异常检测", "近 14 天 vs 历史基线（相对偏离 + 中位数/MAD 稳健 z-score）"],
        ["严重度分级", "高：偏离 ≥100% 或稳健 z ≥6；中：偏离 ≥50% 或 z ≥3.5；其余为低"],
        ["生成时间", evidence["meta"]["generated_at"]],
        ["提示词版本", evidence["meta"].get("prompt_version", "—")],
    ]
    for row in notes:
        worksheet.append(row)
    _style_header(worksheet, 2)
    _autosize(worksheet, max_width=70)


def build_workbook(evidence: dict[str, Any]):
    """构建多 Sheet 工作簿对象（需安装 openpyxl）。"""
    if not OPENPYXL_AVAILABLE:
        raise ExcelUnavailable("未安装 openpyxl，无法导出 Excel。请执行：pip install openpyxl")
    workbook = Workbook()
    _overview_sheet(workbook, evidence)
    for dimension, label in (("站点", "分站点"), ("品类", "分品类"), ("城市", "分城市")):
        rows = evidence["dimensions"].get(dimension) or []
        if rows:
            _dimension_sheet(workbook, label, dimension, rows, ("dispatch_p90", "arrive_ontime_rate", "first_fix_rate", "complaint_rate", "gross_margin"))
    _trend_sheet(workbook, evidence)
    _anomaly_sheet(workbook, evidence)
    _rule_detail_sheet(workbook, evidence)
    _feedback_sheet(workbook, evidence)
    _dispatch_sheet(workbook, evidence)
    _notes_sheet(workbook, evidence)
    return workbook


def write_workbook(evidence: dict[str, Any], path: str) -> str:
    """写出 xlsx 文件，返回路径。"""
    workbook = build_workbook(evidence)
    workbook.save(path)
    return path
